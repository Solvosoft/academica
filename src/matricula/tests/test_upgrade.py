"""Pruebas de las piezas que cambiaron con la migración a Django 6 / djgentelella 0.6."""
from io import BytesIO

from django.contrib.auth.models import User
from django.template import Context, Template
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils.timezone import now
from openpyxl import load_workbook

from djgentelella.async_notification.models import EmailNotification, EmailTemplate
from djgentelella.async_notification.registry import get_context_config
from djgentelella.async_notification.sending import send_email_from_template

from matricula.models import Category, Course, Enroll, Group, Period, Student, WaitingList
from matricula.tests.utils import login_report_admin
from matricula.utils import EMAIL_TEMPLATES, load_email_templates, organization_names
from membership_core.models import Country, SystemCurrency


def create_enroll(organization='[{"value": "UCR"}]'):
    category = Category.objects.create(name='cat', description='cat')
    course = Course.objects.create(category=category, name='Curso', content='x')
    period = Period.objects.create(name='p1', start_date=now(), finish_date=now())
    group = Group.objects.create(
        period=period, course=course, name='g1', schedule='x', pre_enroll_start=now(),
        pre_enroll_finish=now(), enroll_start=now(), enroll_finish=now(), cost=10,
        maximum=20, is_open=True, flow=0, currency=SystemCurrency.objects.get(currency='USD'))
    user = User.objects.create_user('est', 'est@example.com', 'x', first_name='Ana', last_name='Mora')
    student = Student.objects.create(
        user=user, organization=organization, country=Country.objects.get(code='CR'),
        city='SJ', phone_number='8888', expired_at=now())
    return Enroll.objects.create(group=group, student=student, enroll_finished=True,
                                 enroll_activate=True, course_status='approved')


class EmailTemplatesTestCase(TestCase):
    def test_migration_creates_all_templates(self):
        codes = {code for code, *_rest in EMAIL_TEMPLATES}
        self.assertEqual(set(EmailTemplate.objects.values_list('code', flat=True)), codes)

    def test_contexts_are_registered(self):
        for code, *_rest in EMAIL_TEMPLATES:
            self.assertIsNotNone(get_context_config(code), code)

    def test_load_is_idempotent(self):
        EmailTemplate.objects.filter(code='email_open_group').update(subject='Editado')
        self.assertEqual(load_email_templates(), [])
        self.assertEqual(EmailTemplate.objects.get(code='email_open_group').subject, 'Editado')
        self.assertIn('email_open_group', load_email_templates(overwrite=True))

    def test_every_template_renders(self):
        enroll = create_enroll()
        context = Context({'group': enroll.group, 'student': enroll.student, 'user': enroll.student.user,
                           'url': 'http://x', 'domain': 'http://localhost', 'hours_to_pay': 4})
        for template in EmailTemplate.objects.all():
            Template(template.message).render(context)

    def test_send_email_from_template(self):
        enroll = create_enroll()
        notification = send_email_from_template(
            'email_open_group', 'est@example.com',
            {'group': enroll.group, 'url': 'http://x', 'domain': 'http://localhost'}, enqueued=True)
        self.assertEqual(notification.status, 'pending')
        self.assertEqual(notification.recipients, ['est@example.com'])
        self.assertIn(enroll.group.name, notification.message)


class ExportsTestCase(TestCase):
    def setUp(self):
        self.enroll = create_enroll()
        WaitingList.objects.create(group=self.enroll.group, student=self.enroll.student)
        login_report_admin(self.client)

    def test_export_enrolled_group_pdf(self):
        response = self.client.get(reverse('export_enrolled_group', args=[self.enroll.group.pk]),
                                   {'finished': 1})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertTrue(response.content.startswith(b'%PDF'))

    def test_export_waitinglist_xlsx(self):
        response = self.client.get(reverse('export_waitinglist_group', args=[self.enroll.group.pk]))
        rows = list(load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True))
        self.assertEqual(rows[0][0], 'Grupo')
        self.assertEqual(rows[1][:6], ('g1', 'est', 'est@example.com', 'Ana', 'Mora', 'UCR'))

    def test_export_student_status_xlsx(self):
        response = self.client.get(reverse('export_student_status_xls',
                                           args=[self.enroll.group.period.pk, 'approved']))
        rows = list(load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True))
        self.assertEqual(rows[1][-1], 'Curso Aprobado')


class AjaxTestCase(TestCase):
    def setUp(self):
        self.enroll = create_enroll()
        self.client.force_login(self.enroll.student.user)

    def test_requires_ajax_header(self):
        url = reverse('addmetoquee', args=[self.enroll.group.pk])
        self.assertEqual(self.client.get(url).status_code, 400)

    def test_envelope(self):
        url = reverse('addmetoquee', args=[self.enroll.group.pk])
        response = self.client.get(url, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(response.json()['status'], 200)
        self.assertEqual(response.json()['statusText'], 'OK')

    def test_recover_password_fragment(self):
        self.client.logout()
        response = self.client.post(reverse('mail_recover_pass'), {'email': 'est@example.com'},
                                    HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertIn('#recover_pass', response.json()['content']['inner-fragments'])
        self.assertTrue(EmailNotification.objects.filter(recipients=['est@example.com']).exists())


class OrganizationNamesTestCase(TestCase):
    def test_formats(self):
        self.assertEqual(organization_names('[{"value": "UCR"}, {"value": "TEC"}]'), ['UCR', 'TEC'])
        self.assertEqual(organization_names('UCR'), ['UCR'])
        self.assertEqual(organization_names('"UCR"'), ['"UCR"'])
        self.assertEqual(organization_names(''), [])
